"""
업종 유니버스 추출 스크립트
엑셀 파일의 '업종비중' 시트에서 업종코드→종목코드 매핑을 추출합니다.

사용법:
    python3 extract_sector.py <엑셀파일경로>
    예) python3 extract_sector.py ~/Downloads/sector_data.xlsm

출력: sector_universe.json
"""

import sys
import json
import re
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("openpyxl 설치 필요: pip3 install openpyxl")
    sys.exit(1)

SECTOR_FILE = Path(__file__).parent / 'sector_universe.json'


def is_ticker(v: str) -> bool:
    """6자리 숫자 코드인지 확인"""
    return bool(re.fullmatch(r'\d{6}', str(v).strip()))


def is_sector_code(v: str) -> bool:
    """WI/G 업종코드인지 확인"""
    s = str(v).strip()
    return bool(re.match(r'^(WI|G)\d+', s))


def load_name_map(wb) -> dict:
    """관심도db 시트에서 업종코드→업종명 매핑 로드"""
    name_map = {}
    for sheet_name in wb.sheetnames:
        if '관심도' in sheet_name or '관심' in sheet_name:
            ws = wb[sheet_name]
            for row in ws.iter_rows(values_only=True):
                if row and len(row) >= 2:
                    code = str(row[0] or '').strip()
                    name = str(row[1] or '').strip()
                    if is_sector_code(code) and name:
                        name_map[code] = name
            print(f"  관심도db 시트 '{sheet_name}': {len(name_map)}개 업종명 로드")
            break
    return name_map


def parse_sector_sheet(ws, name_map: dict) -> dict:
    """
    '업종비중' 시트 파싱.
    지원 형식:
      형식 A) 행마다 (업종코드, ..., 종목코드, ...) - 업종코드 열과 종목코드 열이 혼재
      형식 B) 업종코드 행 다음에 종목코드 행들이 이어지는 계층 구조
    """
    sectors = {}  # code -> {name, tickers:[]}
    rows = list(ws.iter_rows(values_only=True))

    print(f"\n  시트 행수: {len(rows)}, 열수: {ws.max_column}")
    print("  첫 5행 샘플:")
    for r in rows[:5]:
        print(f"    {[str(v)[:15] if v is not None else 'None' for v in r[:8]]}")

    # --- 형식 A: 각 행에 업종코드 + 종목코드들이 있는 형태 ---
    # 열 A = 업종코드, 열 B = 업종명, 열 C+ = 종목코드들
    format_a_count = 0
    for row in rows[1:]:  # 헤더 스킵
        if not row or row[0] is None:
            continue
        col0 = str(row[0]).strip()
        if not is_sector_code(col0):
            continue
        sec_code = col0
        sec_name = name_map.get(sec_code, '')
        if not sec_name and len(row) > 1 and row[1]:
            sec_name = str(row[1]).strip()
        tickers = []
        for cell in row[2:]:
            if cell and is_ticker(str(cell)):
                tickers.append(str(cell).strip().zfill(6))
        if tickers:
            format_a_count += 1
            sectors[sec_code] = {'name': sec_name or sec_code, 'tickers': tickers}

    if format_a_count > 0:
        print(f"\n  형식 A 감지: {format_a_count}개 업종 (업종코드 + 종목코드 행)")
        return sectors

    # --- 형식 B: 업종코드 행 → 이어지는 종목코드 행들 ---
    current_sector = None
    for row in rows:
        if not row or row[0] is None:
            continue
        col0 = str(row[0]).strip()
        if is_sector_code(col0):
            sec_code = col0
            sec_name = name_map.get(sec_code, '')
            if not sec_name and len(row) > 1 and row[1]:
                sec_name = str(row[1]).strip()
            current_sector = sec_code
            if sec_code not in sectors:
                sectors[sec_code] = {'name': sec_name or sec_code, 'tickers': []}
        elif current_sector:
            for cell in row:
                if cell and is_ticker(str(cell)):
                    sectors[current_sector]['tickers'].append(str(cell).strip().zfill(6))

    if sectors:
        print(f"\n  형식 B 감지: {len(sectors)}개 업종 (계층 구조)")
        return sectors

    # --- 형식 C: 컬럼 탐색 ---
    # 업종코드 컬럼과 종목코드 컬럼을 자동 탐색
    print("\n  형식 A/B 미감지, 컬럼 자동 탐색...")
    sector_col = ticker_col = name_col = None
    header = rows[0] if rows else []
    for i, h in enumerate(header):
        if h is None: continue
        hs = str(h).lower()
        if any(k in hs for k in ['업종코드', '섹터코드', 'sector_code', 'code']):
            sector_col = i
        elif any(k in hs for k in ['종목코드', '코드', 'ticker', 'stock']):
            ticker_col = i
        elif any(k in hs for k in ['업종명', '섹터명', 'name', '명']):
            name_col = i

    if sector_col is not None and ticker_col is not None:
        for row in rows[1:]:
            if not row or len(row) <= max(sector_col, ticker_col):
                continue
            sec = str(row[sector_col] or '').strip()
            tck = str(row[ticker_col] or '').strip()
            if is_sector_code(sec) and is_ticker(tck):
                sec_name = name_map.get(sec, '')
                if not sec_name and name_col is not None and len(row) > name_col:
                    sec_name = str(row[name_col] or '').strip()
                if sec not in sectors:
                    sectors[sec] = {'name': sec_name or sec, 'tickers': []}
                sectors[sec]['tickers'].append(tck.zfill(6))
        print(f"  형식 C 감지: {len(sectors)}개 업종 (컬럼 기반)")
        return sectors

    return sectors


def main():
    if len(sys.argv) < 2:
        print("사용법: python3 extract_sector.py <엑셀파일경로>")
        sys.exit(1)

    excel_path = Path(sys.argv[1]).expanduser()
    if not excel_path.exists():
        print(f"파일 없음: {excel_path}")
        sys.exit(1)

    print(f"엑셀 파일 로드: {excel_path}")
    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    print(f"시트 목록: {wb.sheetnames}")

    # 관심도db 시트에서 업종명 로드
    name_map = load_name_map(wb)

    # 업종비중 시트 찾기
    target_sheet = None
    for name in wb.sheetnames:
        if '업종비중' in name or '업종' in name:
            target_sheet = name
            break
    if target_sheet is None:
        print(f"\n'업종비중' 시트를 찾을 수 없습니다.")
        print(f"사용 가능한 시트: {wb.sheetnames}")
        sys.exit(1)

    print(f"\n파싱 시트: '{target_sheet}'")
    ws = wb[target_sheet]
    sectors = parse_sector_sheet(ws, name_map)

    if not sectors:
        print("\n업종 데이터를 파싱하지 못했습니다.")
        print("시트 구조를 확인하세요.")
        sys.exit(1)

    # 이름 없는 업종에 name_map 적용
    for code, info in sectors.items():
        if not info['name'] or info['name'] == code:
            info['name'] = name_map.get(code, code)

    # 통계
    total_tickers = sum(len(v['tickers']) for v in sectors.values())
    print(f"\n결과: {len(sectors)}개 업종, {total_tickers}개 종목 매핑")
    for code, info in list(sectors.items())[:5]:
        print(f"  {code}: {info['name']} ({len(info['tickers'])}종목) "
              f"- {info['tickers'][:3]}...")

    SECTOR_FILE.write_text(json.dumps(sectors, ensure_ascii=False, indent=2))
    print(f"\n저장 완료: {SECTOR_FILE}")


if __name__ == '__main__':
    main()
